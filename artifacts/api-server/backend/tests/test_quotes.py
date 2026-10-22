"""Tests for the Quote agent routes — chat (start + continue), edit, export, tenancy.

Fully offline: the Anthropic key check is forced "configured" and the unified
chat_turn function is monkeypatched, so no network or real credentials run.
Routes are called directly (the suite's convention).
"""
from __future__ import annotations

import asyncio
import io
import json

import pytest
from fastapi import HTTPException
from starlette.datastructures import Headers, UploadFile

from app.models import Company, Conversation, Organization, Quote, User, Visitor
from app.routers import quotes as router
from app.services import document_parse
from app.services import quote_export as qx


def _upload(filename: str, data: bytes) -> UploadFile:
    """Build a Starlette UploadFile the way the multipart endpoint receives it."""
    return UploadFile(
        file=io.BytesIO(data),
        filename=filename,
        headers=Headers({"content-type": "application/octet-stream"}),
    )


# --- fixtures / helpers ---------------------------------------------------

def _user(db, org, email="rep@example.com", role="owner") -> User:
    u = User(org_id=org.id, email=email, name="Rep", password_hash="x", role=role)
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _other_org_user(db) -> User:
    o = Organization(name="Other Org", org_type="company")
    db.add(o)
    db.commit()
    db.refresh(o)
    return _user(db, o, email="intruder@example.com")


def _visitor(db, org, name="Jane Doe", company="Acme", company_id=None) -> Visitor:
    v = Visitor(org_id=org.id, name=name, company=company, company_id=company_id)
    db.add(v)
    db.commit()
    db.refresh(v)
    return v


def _company(db, org, name="Acme") -> Company:
    c = Company(org_id=org.id, name=name)
    db.add(c)
    db.commit()
    db.refresh(c)
    return c


def _conversation(db, org, visitor, company_id=None, transcript="We need 100 widgets next quarter.") -> Conversation:
    conv = Conversation(
        org_id=org.id,
        visitor_id=visitor.id,
        company_id=company_id,
        transcript=transcript,
        topic="Widget supply",
        status="analyzed",
    )
    db.add(conv)
    db.commit()
    db.refresh(conv)
    return conv


@pytest.fixture
def ai_ok(monkeypatch):
    """Force the AI 'configured' check on and replace chat_turn with a
    deterministic stub: the first turn drafts two line items; later turns append
    a support retainer and tag the narrative so refinement is observable."""
    monkeypatch.setattr(router.quote_ai, "describe_anthropic_problem", lambda: None)

    def _turn(context, current, history, message):
        if current is None:
            return {
                "reply": "Here's a first draft based on your conversation(s).",
                "title": "Widget Supply Proposal",
                "client_name": context.get("client_name"),
                "client_company": context.get("client_company"),
                "proposal_text": "Thanks for the conversation. Here's our proposal.",
                "line_items": [
                    {"description": "Widgets", "quantity": 100, "unit_price": 5.0},
                    {"description": "Setup fee", "quantity": 1, "unit_price": 250.0},
                ],
                "notes": "Valid for 30 days.",
            }
        items = list(current.get("line_items") or [])
        items.append({"description": "Support retainer", "quantity": 12, "unit_price": 100.0})
        return {
            "reply": "Added a 12-month support retainer.",
            "title": current.get("title"),
            "client_name": None,
            "client_company": None,
            "proposal_text": (current.get("proposal_text") or "") + " (revised)",
            "line_items": items,
            "notes": current.get("notes"),
        }

    monkeypatch.setattr(router.quote_ai, "chat_turn", _turn)


# --- chat: start a new proposal ------------------------------------------

def test_chat_start_from_conversation(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    out = router.chat(
        router.ChatIn(message="Draft a proposal", conversation_ids=[conv.id]),
        db=db, user=user,
    )
    assert out["source_type"] == "conversation"
    assert out["source_conversation_id"] == conv.id
    assert out["source_conversation_ids"] == [conv.id]
    assert out["title"] == "Widget Supply Proposal"
    assert len(out["line_items"]) == 2
    # subtotal = 100*5 + 1*250 = 750
    assert out["totals"]["subtotal"] == 750.0
    assert out["totals"]["total"] == 750.0
    # The thread seeds both the user message and the agent reply.
    assert [m["role"] for m in out["chat_messages"]] == ["user", "assistant"]
    assert out["chat_messages"][0]["text"] == "Draft a proposal"


def test_chat_start_multiple_conversations(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    c1 = _conversation(db, org, visitor, transcript="First chat about widgets.")
    c2 = _conversation(db, org, visitor, transcript="Follow up on volume pricing.")
    out = router.chat(
        router.ChatIn(message="Combine these", conversation_ids=[c1.id, c2.id]),
        db=db, user=user,
    )
    assert out["source_conversation_ids"] == [c1.id, c2.id]
    assert out["source_type"] == "conversation"


def test_chat_start_from_contact_aggregates(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    _conversation(db, org, visitor, transcript="First chat about widgets.")
    _conversation(db, org, visitor, transcript="Follow up on volume pricing.")
    out = router.chat(
        router.ChatIn(message="Draft from everything we discussed", visitor_id=visitor.id),
        db=db, user=user,
    )
    assert out["source_type"] == "contact"
    assert out["visitor_id"] == visitor.id
    assert out["client_company"] == "Acme"
    # Both of the contact's conversations were aggregated as the grounding set.
    assert len(out["source_conversation_ids"]) == 2


def test_chat_start_unions_explicit_convs_with_pinned_contact(db, org, ai_ok):
    # Explicit conversation chip PLUS a pinned contact must union: the standalone
    # conversation and both of the contact's conversations, de-duped.
    user = _user(db, org)
    contact = _visitor(db, org, name="Pinned Person")
    a = _conversation(db, org, contact, transcript="Contact chat A.")
    b = _conversation(db, org, contact, transcript="Contact chat B.")
    other_visitor = _visitor(db, org, name="Someone Else", company="ZCo")
    standalone = _conversation(db, org, other_visitor, transcript="Standalone chat.")
    out = router.chat(
        router.ChatIn(
            message="Combine these",
            conversation_ids=[standalone.id],
            visitor_id=contact.id,
        ),
        db=db, user=user,
    )
    assert set(out["source_conversation_ids"]) == {standalone.id, a.id, b.id}
    # Explicit chip is listed first; contact's conversations follow.
    assert out["source_conversation_ids"][0] == standalone.id


def test_chat_start_from_company(db, org, ai_ok):
    user = _user(db, org)
    company = _company(db, org)
    visitor = _visitor(db, org, company_id=company.id)
    _conversation(db, org, visitor, company_id=company.id)
    out = router.chat(
        router.ChatIn(message="Draft for this company", company_id=company.id),
        db=db, user=user,
    )
    assert out["source_type"] == "company"
    assert out["company_id"] == company.id
    assert out["client_company"] == "Acme"


def test_chat_blocked_without_ai(db, org, monkeypatch):
    monkeypatch.setattr(router.quote_ai, "describe_anthropic_problem", lambda: "AI not configured")
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    with pytest.raises(HTTPException) as ei:
        router.chat(
            router.ChatIn(message="Draft it", conversation_ids=[conv.id]),
            db=db, user=user,
        )
    assert ei.value.status_code == 400


def test_chat_requires_message(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    with pytest.raises(HTTPException) as ei:
        router.chat(
            router.ChatIn(message="   ", conversation_ids=[conv.id]),
            db=db, user=user,
        )
    assert ei.value.status_code == 400


def test_chat_no_transcript_rejected(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    _conversation(db, org, visitor, transcript="   ")
    with pytest.raises(HTTPException) as ei:
        router.chat(
            router.ChatIn(message="Draft it", visitor_id=visitor.id),
            db=db, user=user,
        )
    assert ei.value.status_code == 400


def test_chat_does_not_leak_cross_org_company(db, org, ai_ok):
    # A conversation in `org` carries a dangling company_id pointing at another
    # org's Company. The drafted quote must NOT pick up that company's name.
    other = Organization(name="Rival Corp", org_type="company")
    db.add(other)
    db.commit()
    db.refresh(other)
    foreign = Company(org_id=other.id, name="Rival Secret Co")
    db.add(foreign)
    db.commit()
    db.refresh(foreign)

    user = _user(db, org)
    visitor = _visitor(db, org, company="Our Client")
    conv = _conversation(db, org, visitor, company_id=foreign.id)
    out = router.chat(
        router.ChatIn(message="Draft it", conversation_ids=[conv.id]),
        db=db, user=user,
    )
    assert out["client_company"] != "Rival Secret Co"
    assert out["client_company"] == "Our Client"


def test_chat_cross_org_conversation_404(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    intruder = _other_org_user(db)
    with pytest.raises(HTTPException) as ei:
        router.chat(
            router.ChatIn(message="Draft it", conversation_ids=[conv.id]),
            db=db, user=intruder,
        )
    assert ei.value.status_code == 404


# --- chat: continue (refine + thread + versions) -------------------------

def _make_quote(db, org, user) -> dict:
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    return router.chat(
        router.ChatIn(message="Draft a proposal", conversation_ids=[conv.id]),
        db=db, user=user,
    )


def test_chat_continue_refines_threads_and_snapshots(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    before_items = len(q["line_items"])
    out = router.chat(
        router.ChatIn(quote_id=q["id"], message="Add a support retainer"),
        db=db, user=user,
    )
    # Refinement applied.
    assert len(out["line_items"]) == before_items + 1
    assert "(revised)" in out["proposal_text"]
    # Prior state snapshotted for undo.
    assert out["version_count"] == 1
    assert out["versions"][0]["feedback"] == "Add a support retainer"
    # Thread now has the seed pair plus this turn's pair.
    assert len(out["chat_messages"]) == 4
    assert out["chat_messages"][-2]["text"] == "Add a support retainer"
    assert out["chat_messages"][-1]["role"] == "assistant"


def test_chat_continue_cross_org_404(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    intruder = _other_org_user(db)
    with pytest.raises(HTTPException) as ei:
        router.chat(
            router.ChatIn(quote_id=q["id"], message="Hijack"),
            db=db, user=intruder,
        )
    assert ei.value.status_code == 404


def test_chat_continue_merges_attached_context(db, org, ai_ok):
    # A follow-up turn can bring in NEW source context with the message; it must
    # merge into the quote's grounding set and persist.
    user = _user(db, org)
    q = _make_quote(db, org, user)
    orig = q["source_conversation_ids"][0]
    v2 = _visitor(db, org, name="Bob", company="Globex")
    conv2 = _conversation(db, org, v2)
    out = router.chat(
        router.ChatIn(quote_id=q["id"], message="Also factor in this call",
                      conversation_ids=[conv2.id]),
        db=db, user=user,
    )
    assert set(out["source_conversation_ids"]) == {orig, conv2.id}


def test_chat_continue_preserves_proposal_on_empty_response(db, org, ai_ok, monkeypatch):
    # A chat-only / hiccup response (no proposal body) must not wipe the proposal.
    user = _user(db, org)
    q = _make_quote(db, org, user)
    monkeypatch.setattr(
        router.quote_ai, "chat_turn",
        lambda context, current, history, message: {
            "reply": "Sure, what would you like to change?",
            "title": "", "client_name": None, "client_company": None,
            "proposal_text": "", "line_items": [], "notes": None,
        },
    )
    out = router.chat(
        router.ChatIn(quote_id=q["id"], message="hmm"), db=db, user=user,
    )
    assert out["proposal_text"] == q["proposal_text"]
    assert out["line_items"] == q["line_items"]
    assert out["chat_messages"][-1]["text"] == "Sure, what would you like to change?"


def test_chat_continue_allows_intentional_clear(db, org, ai_ok, monkeypatch):
    # When the model re-renders the full proposal with no line items (e.g. the
    # user asked to remove them), the empty list IS applied.
    user = _user(db, org)
    q = _make_quote(db, org, user)
    assert len(q["line_items"]) == 2
    monkeypatch.setattr(
        router.quote_ai, "chat_turn",
        lambda context, current, history, message: {
            "reply": "Removed all line items.",
            "title": "Widget Supply Proposal",
            "client_name": None, "client_company": None,
            "proposal_text": "Updated proposal with no priced items.",
            "line_items": [], "notes": None,
        },
    )
    out = router.chat(
        router.ChatIn(quote_id=q["id"], message="remove all line items"),
        db=db, user=user,
    )
    assert out["line_items"] == []
    assert out["proposal_text"] == "Updated proposal with no priced items."


# --- chat: context endpoint (in-thread add/remove) -----------------------

def test_context_get_reflects_source(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    orig = q["source_conversation_ids"][0]
    ctx = router.get_context(q["id"], db=db, user=user)
    assert [c["id"] for c in ctx["conversations"]] == [orig]
    # The conversation's contact is inferred at creation and surfaces as a chip.
    assert ctx["visitor"] is not None
    assert ctx["company"] is None


def test_context_set_adds_and_removes_conversations(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    qid, orig = q["id"], q["source_conversation_ids"][0]
    v2 = _visitor(db, org, name="Bob", company="Globex")
    conv2 = _conversation(db, org, v2)
    # Add a second conversation.
    ctx = router.set_context(
        qid, router.QuoteContextIn(conversation_ids=[orig, conv2.id]),
        db=db, user=user,
    )
    assert set(c["id"] for c in ctx["conversations"]) == {orig, conv2.id}
    # Remove the original (full explicit set without it) — it stays removed.
    ctx = router.set_context(
        qid, router.QuoteContextIn(conversation_ids=[conv2.id]),
        db=db, user=user,
    )
    assert [c["id"] for c in ctx["conversations"]] == [conv2.id]


def test_context_pinned_contact_expands_then_individual_removal_sticks(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    qid, orig = q["id"], q["source_conversation_ids"][0]
    contact = _visitor(db, org, name="Carol", company="Initech")
    a = _conversation(db, org, contact)
    b = _conversation(db, org, contact)
    # Pinning the contact expands its conversations into the set.
    ctx = router.set_context(
        qid, router.QuoteContextIn(conversation_ids=[orig], visitor_id=contact.id),
        db=db, user=user,
    )
    assert set(c["id"] for c in ctx["conversations"]) == {orig, a.id, b.id}
    assert ctx["visitor"]["id"] == contact.id
    # Remove one expanded conversation while the contact stays pinned: no
    # re-expansion, so the removal sticks.
    ctx = router.set_context(
        qid, router.QuoteContextIn(conversation_ids=[orig, b.id], visitor_id=contact.id),
        db=db, user=user,
    )
    ids = set(c["id"] for c in ctx["conversations"])
    assert a.id not in ids and ids == {orig, b.id}


def test_context_set_grounds_next_chat_turn(db, org, ai_ok):
    # Context changes must drive the next chat turn's grounding.
    user = _user(db, org)
    q = _make_quote(db, org, user)
    qid = q["id"]
    v2 = _visitor(db, org, name="Bob", company="Globex")
    conv2 = _conversation(db, org, v2)
    router.set_context(
        qid, router.QuoteContextIn(conversation_ids=[conv2.id]),
        db=db, user=user,
    )
    out = router.chat(
        router.ChatIn(quote_id=qid, message="refine"), db=db, user=user,
    )
    assert out["source_conversation_ids"] == [conv2.id]


def test_context_get_cross_org_404(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    intruder = _other_org_user(db)
    with pytest.raises(HTTPException) as ei:
        router.get_context(q["id"], db=db, user=intruder)
    assert ei.value.status_code == 404


def test_context_rejects_cross_org_conversation(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    intruder = _other_org_user(db)
    other_org = db.get(Organization, intruder.org_id)
    foreign_visitor = _visitor(db, other_org, name="Eve", company="Hooli")
    foreign_conv = _conversation(db, other_org, foreign_visitor)
    with pytest.raises(HTTPException) as ei:
        router.set_context(
            q["id"],
            router.QuoteContextIn(conversation_ids=[foreign_conv.id]),
            db=db, user=user,
        )
    assert ei.value.status_code == 404


# --- update / edit --------------------------------------------------------

def test_update_recomputes_totals(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    out = router.update_quote(
        q["id"],
        router.QuoteUpdateIn(
            line_items=[router.LineItemIn(description="Widgets", quantity=10, unit_price=5)],
            tax_percent=10,
            discount_percent=0,
        ),
        db=db, user=user,
    )
    # subtotal 50, tax 10% => total 55
    assert out["totals"]["subtotal"] == 50.0
    assert out["totals"]["tax_amount"] == 5.0
    assert out["totals"]["total"] == 55.0


def test_update_drops_blank_line_items(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    out = router.update_quote(
        q["id"],
        router.QuoteUpdateIn(line_items=[
            router.LineItemIn(description="Real", quantity=1, unit_price=10),
            router.LineItemIn(description="   ", quantity=2, unit_price=5),
        ]),
        db=db, user=user,
    )
    assert len(out["line_items"]) == 1


def test_update_cross_org_404(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    intruder = _other_org_user(db)
    with pytest.raises(HTTPException) as ei:
        router.update_quote(q["id"], router.QuoteUpdateIn(title="Hijacked"), db=db, user=intruder)
    assert ei.value.status_code == 404


# --- delete / list --------------------------------------------------------

def test_delete_quote(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    router.delete_quote(q["id"], db=db, user=user)
    assert db.get(Quote, q["id"]) is None


def test_list_scoped_to_org(db, org, ai_ok):
    user = _user(db, org)
    _make_quote(db, org, user)
    intruder = _other_org_user(db)
    assert router.list_quotes(limit=100, brand_id=None, db=db, user=user) != []
    assert router.list_quotes(limit=100, brand_id=None, db=db, user=intruder) == []


# --- export ---------------------------------------------------------------

def test_export_pdf_and_docx(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    pdf = router.export_pdf(q["id"], db=db, user=user)
    docx = router.export_docx(q["id"], db=db, user=user)
    # StreamingResponse — confirm the media types are correct.
    assert pdf.media_type == "application/pdf"
    assert docx.media_type.endswith("wordprocessingml.document")


def test_export_cross_org_404(db, org, ai_ok):
    user = _user(db, org)
    q = _make_quote(db, org, user)
    intruder = _other_org_user(db)
    with pytest.raises(HTTPException) as ei:
        router.export_pdf(q["id"], db=db, user=intruder)
    assert ei.value.status_code == 404


# --- export rendering (quote_export) -------------------------------------

_BRANDING = {
    "company_name": "Café Naïve Marine",
    "address": "123 Harbour Rd",
    "email": "sales@example.com",
    "phone": "+1 (555) 010-9999",
    "website": "example.com",
    "primary_color": "#0E2A47",
    "accent_color": "#FF5B2E",
    "logo_data_url": None,
}

_LONG_DESC = (
    "Full custom fit-out of the flybridge including teak decking, hand-stitched "
    "leather helm seating, integrated NMEA 2000 sensor network, and a "
    "\u201cconcierge\u201d commissioning package \u2014 delivered dockside."
)


def test_unicode_font_registers():
    # The bundled DejaVu font must register so exports keep full Unicode (no
    # lossy ASCII fallback in the common case).
    from fpdf import FPDF

    assert qx._register_unicode_font(FPDF()) is True


def test_render_pdf_preserves_unicode_glyphs():
    # Smart quotes / em-dash / accented names render via the embedded font, and
    # the font subset is actually embedded in the output.
    quote = {
        "title": "\u201cMy Dream Boat\u201d \u2014 SOW",
        "client_name": "Jos\u00e9 M\u00fcller",
        "client_company": "LPGA Marine, Inc.",
        "currency": "\u20ac",
        "proposal_text": "Grounded in our conversations \u2014 here is the scope.",
        "line_items": [{"description": _LONG_DESC, "quantity": 1, "unit_price": 48500}],
        "tax_percent": 8.25,
        "discount_percent": 5,
        "notes": "Valid 30 days \u2014 50% deposit on signing.",
    }
    pdf = qx.render_pdf(quote, _BRANDING)
    assert pdf[:4] == b"%PDF"
    assert b"DejaVu" in pdf  # embedded subset proves no ASCII downgrade


@pytest.mark.parametrize(
    "quote",
    [
        # No line items at all.
        {"title": "Discovery", "proposal_text": "Narrative only.", "line_items": []},
        # Many wrapping rows that force a page break (table must not overlap).
        {
            "title": "Bulk Order",
            "client_company": "Globex",
            "line_items": [
                {
                    "description": f"Widget {i} with a long descriptive name that wraps",
                    "quantity": i,
                    "unit_price": i * 1.5,
                }
                for i in range(1, 41)
            ],
            "tax_percent": 10,
        },
        # Degenerate / odd values: missing fields, non-numeric qty, odd currency.
        {
            "title": "",
            "client_name": None,
            "client_company": None,
            "proposal_text": None,
            "currency": "\u00a5",
            "line_items": [{"description": "", "quantity": "x", "unit_price": None}],
            "notes": "",
        },
    ],
)
def test_render_pdf_and_docx_never_crash(quote):
    pdf = qx.render_pdf(quote, _BRANDING)
    docx = qx.render_docx(quote, _BRANDING)
    assert pdf[:4] == b"%PDF"
    assert docx[:2] == b"PK"


def test_render_handles_logo_and_missing_logo():
    import base64

    one_px_png = base64.b64encode(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8z8BQ"
        "DwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )).decode()
    with_logo = {**_BRANDING, "logo_data_url": "data:image/png;base64," + one_px_png}
    bad_logo = {**_BRANDING, "logo_data_url": "data:image/png;base64,not-valid"}
    quote = {"title": "Logo Test", "line_items": [{"description": "X", "quantity": 1, "unit_price": 5}]}
    for branding in (with_logo, bad_logo):
        assert qx.render_pdf(quote, branding)[:4] == b"%PDF"
        assert qx.render_docx(quote, branding)[:2] == b"PK"


def test_ascii_fallback_when_font_missing(monkeypatch):
    # If the bundled font can't be registered, export must still succeed using
    # the latin-1 core font (smart punctuation sanitised, no crash).
    monkeypatch.setattr(qx, "_register_unicode_font", lambda pdf: False)
    quote = {
        "title": "Sm\u00f6rg\u00e5sbord \u2014 \u201cPlan\u201d",
        "line_items": [{"description": "Caf\u00e9 \u2014 item", "quantity": 1, "unit_price": 9}],
        "notes": "Don\u2019t panic.",
    }
    pdf = qx.render_pdf(quote, _BRANDING)
    assert pdf[:4] == b"%PDF"


# --- rich-text formatting (bullets / headings / bold) --------------------

def test_parse_rich_blocks_groups_bullets_headings_and_prose():
    text = (
        "Intro paragraph about the client.\n"
        "\n"
        "## Scope\n"
        "- First deliverable\n"
        "* Second deliverable\n"
        "\u2022 Third deliverable\n"
        "\n"
        "Closing **next step** paragraph."
    )
    blocks = qx.parse_rich_blocks(text)
    types = [b["type"] for b in blocks]
    assert types == ["para", "heading", "bullets", "para"]
    assert blocks[1]["text"] == "Scope"
    # All three bullet glyphs (-, *, •) collapse into one grouped list.
    assert blocks[2]["items"] == ["First deliverable", "Second deliverable", "Third deliverable"]
    assert blocks[3]["text"] == "Closing **next step** paragraph."


def test_parse_rich_blocks_tolerant_of_empty_and_loose_input():
    assert qx.parse_rich_blocks("") == []
    assert qx.parse_rich_blocks(None) == []
    # A bare "## " heading marker with no text is dropped, not emitted empty.
    assert qx.parse_rich_blocks("## \n- \nReal line") == [{"type": "para", "text": "Real line"}]


def test_inline_runs_splits_bold_and_keeps_unmatched_marker():
    assert qx._inline_runs("plain") == [("plain", False)]
    assert qx._inline_runs("a **bold** c") == [("a ", False), ("bold", True), (" c", False)]
    # An unmatched trailing "**" stays literal rather than vanishing.
    assert qx._inline_runs("oops ** here") == [("oops ** here", False)]


def _docx_paragraphs(docx_bytes):
    import io as _io
    from docx import Document
    return Document(_io.BytesIO(docx_bytes)).paragraphs


def test_render_docx_renders_bullets_headings_and_bold():
    quote = {
        "title": "Structured Proposal",
        "proposal_text": (
            "Opening narrative.\n\n"
            "## Deliverables\n"
            "- Teak decking\n"
            "- **Concierge** commissioning\n\n"
            "Closing line."
        ),
        "line_items": [{"description": "Fit-out", "quantity": 1, "unit_price": 1000}],
        "notes": "## Terms\n- Valid 30 days\n- 50% deposit",
    }
    docx = qx.render_docx(quote, _BRANDING)
    assert docx[:2] == b"PK"
    paras = _docx_paragraphs(docx)
    styles = {p.style.name for p in paras}
    assert "List Bullet" in styles
    texts = [p.text for p in paras]
    # No raw bullet glyphs leaked into a List Bullet paragraph's text.
    bullet_texts = [p.text for p in paras if p.style.name == "List Bullet"]
    assert any("Teak decking" in t for t in bullet_texts)
    assert all(not t.lstrip().startswith(("-", "*", "\u2022")) for t in bullet_texts)
    # The sub-heading text survives (markers stripped) and a bold run exists.
    assert any(t == "Deliverables" for t in texts)
    bold_runs = [r.text for p in paras for r in p.runs if r.bold]
    assert "Concierge" in bold_runs


def test_render_pdf_with_rich_formatting_never_crashes():
    quote = {
        "title": "Rich PDF",
        "proposal_text": "Intro.\n\n## Scope\n- One\n- Two with **bold**\n\nClose.",
        "line_items": [{"description": "Item", "quantity": 1, "unit_price": 10}],
        "notes": "## Terms\n- Net 30\n- No refunds",
    }
    pdf = qx.render_pdf(quote, _BRANDING)
    assert pdf[:4] == b"%PDF"


# --- document_parse service ----------------------------------------------

def test_parse_txt_and_md_return_text():
    assert "hello world" in document_parse.parse_upload("a.txt", b"hello world")
    assert "Heading" in document_parse.parse_upload("b.md", b"# Heading\n\nBody.")


def test_parse_docx_returns_text():
    from docx import Document  # python-docx, a project dependency

    doc = Document()
    doc.add_paragraph("Hello from a Word document.")
    buf = io.BytesIO()
    doc.save(buf)
    text = document_parse.parse_upload("brief.docx", buf.getvalue())
    assert "Hello from a Word document." in text


def test_parse_unsupported_extension_rejected():
    with pytest.raises(document_parse.DocumentParseError):
        document_parse.parse_upload("malware.exe", b"MZ binary")


def test_parse_empty_file_rejected():
    with pytest.raises(document_parse.DocumentParseError):
        document_parse.parse_upload("empty.txt", b"")


def test_parse_oversized_file_rejected():
    big = b"x" * (document_parse.MAX_FILE_BYTES + 1)
    with pytest.raises(document_parse.DocumentParseError):
        document_parse.parse_upload("big.txt", big)


def test_parse_textless_file_rejected():
    # Whitespace-only content yields no readable text.
    with pytest.raises(document_parse.DocumentParseError):
        document_parse.parse_upload("blank.txt", b"   \n\t  ")


def test_build_payload_too_many_files_rejected():
    files = [(f"f{i}.txt", b"hi") for i in range(document_parse.MAX_FILES + 1)]
    with pytest.raises(document_parse.DocumentParseError):
        document_parse.build_attachment_payload(files)


def test_merge_attachments_combines_and_handles_empty():
    text, meta = document_parse.merge_attachments(
        None, None, [{"name": "a.txt", "text": "alpha"}]
    )
    assert "alpha" in text
    assert meta == [{"name": "a.txt", "chars": 5}]

    text2, meta2 = document_parse.merge_attachments(
        text, meta, [{"name": "b.txt", "text": "beta"}]
    )
    assert "alpha" in text2 and "beta" in text2
    assert [m["name"] for m in meta2] == ["a.txt", "b.txt"]

    # Nothing in -> nothing out (keeps the columns NULL).
    assert document_parse.merge_attachments(None, None, None) == (None, None)


# --- chat/upload: multipart entry point ----------------------------------

def _capture_chat_turn(monkeypatch, seen):
    def _turn(context, current, history, message):
        seen["context"] = context
        return {
            "reply": "ok",
            "title": (current or {}).get("title") or "Captured",
            "client_name": None,
            "client_company": None,
            "proposal_text": ((current or {}).get("proposal_text") or "") + " x",
            "line_items": (current or {}).get("line_items") or [],
            "notes": None,
        }

    monkeypatch.setattr(router.quote_ai, "chat_turn", _turn)


def test_chat_upload_start_grounds_and_persists(db, org, ai_ok, monkeypatch):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    seen: dict = {}
    _capture_chat_turn(monkeypatch, seen)

    out = asyncio.run(router.chat_upload(
        message="Draft a proposal",
        quote_id=None,
        conversation_ids=json.dumps([conv.id]),
        visitor_id=None,
        company_id=None,
        brand_id=None,
        template_id=None,
        files=[_upload("brief.txt", b"Budget is $5000 for the project.")],
        db=db, user=user,
    ))

    # The attachment text is folded into the grounding context for the draft.
    assert "Budget is $5000" in seen["context"]["attachments_text"]
    # Display metadata is exposed (text itself is not in the API payload).
    assert [m["name"] for m in out["attachments_meta"]] == ["brief.txt"]
    assert "attachments_text" not in out
    # And it is persisted on the quote for later refinements.
    q = db.get(Quote, out["id"])
    assert "Budget is $5000" in q.attachments_text


def test_chat_upload_refine_appends_attachment(db, org, ai_ok, monkeypatch):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)

    # First draft with one attachment (uses ai_ok's default chat_turn).
    out = asyncio.run(router.chat_upload(
        message="Draft",
        quote_id=None,
        conversation_ids=json.dumps([conv.id]),
        visitor_id=None,
        company_id=None,
        brand_id=None,
        template_id=None,
        files=[_upload("a.txt", b"First doc.")],
        db=db, user=user,
    ))
    qid = out["id"]
    assert [m["name"] for m in out["attachments_meta"]] == ["a.txt"]

    # Refinement adds a second attachment; both must reach the grounding context.
    seen: dict = {}
    _capture_chat_turn(monkeypatch, seen)
    out2 = asyncio.run(router.chat_upload(
        message="Refine",
        quote_id=qid,
        conversation_ids=None,
        visitor_id=None,
        company_id=None,
        brand_id=None,
        template_id=None,
        files=[_upload("b.txt", b"Second doc.")],
        db=db, user=user,
    ))
    assert [m["name"] for m in out2["attachments_meta"]] == ["a.txt", "b.txt"]
    assert "First doc." in seen["context"]["attachments_text"]
    assert "Second doc." in seen["context"]["attachments_text"]


def test_chat_upload_without_files_behaves_like_chat(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    out = asyncio.run(router.chat_upload(
        message="Draft",
        quote_id=None,
        conversation_ids=json.dumps([conv.id]),
        visitor_id=None,
        company_id=None,
        brand_id=None,
        template_id=None,
        files=[],
        db=db, user=user,
    ))
    assert out["title"] == "Widget Supply Proposal"
    assert out["attachments_meta"] == []


def test_chat_upload_rejects_bad_file_without_creating_quote(db, org, ai_ok):
    user = _user(db, org)
    visitor = _visitor(db, org)
    conv = _conversation(db, org, visitor)
    with pytest.raises(HTTPException) as ei:
        asyncio.run(router.chat_upload(
            message="Draft",
            quote_id=None,
            conversation_ids=json.dumps([conv.id]),
            visitor_id=None,
            company_id=None,
            brand_id=None,
            template_id=None,
            files=[_upload("malware.exe", b"nope")],
            db=db, user=user,
        ))
    assert ei.value.status_code == 400
    assert db.query(Quote).count() == 0
