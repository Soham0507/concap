"""Render a Quote to a branded PDF (fpdf2) or Word .docx (python-docx).

Both libraries are imported lazily inside the render functions so importing
this module never requires them at import time. The org branding (logo,
colors, contact details) is passed in as a plain dict so the same renderers
work for a live org or a frozen branding snapshot stored on the quote.

The PDF embeds a bundled Unicode font (DejaVu Sans) so smart quotes, em-dashes
and accented client/company names render correctly. If the font asset is ever
missing it falls back to fpdf2's latin-1 core font (with ASCII sanitising) so
export never hard-crashes.
"""
from __future__ import annotations

import base64
import io
import os
import re

# Bundled Unicode font — no system-font assumption (deploy is VM + SQLite).
_FONT_DIR = os.path.join(os.path.dirname(__file__), "..", "assets", "fonts")
_FONT_FAMILY = "DejaVu"
_FONT_FILES = {"": "DejaVuSans.ttf", "B": "DejaVuSans-Bold.ttf"}

_MUTED = (110, 110, 110)
_BODY = (45, 45, 50)


def compute_totals(line_items, tax_percent: float = 0.0, discount_percent: float = 0.0) -> dict:
    """Sum line items and apply discount then tax. Returns
    {subtotal, discount_amount, taxable, tax_amount, total}."""
    subtotal = 0.0
    for it in line_items or []:
        try:
            qty = float(it.get("quantity", 0) or 0)
            price = float(it.get("unit_price", 0) or 0)
        except (TypeError, ValueError, AttributeError):
            continue
        subtotal += qty * price
    discount_amount = subtotal * (float(discount_percent or 0) / 100.0)
    taxable = subtotal - discount_amount
    tax_amount = taxable * (float(tax_percent or 0) / 100.0)
    total = taxable + tax_amount
    return {
        "subtotal": round(subtotal, 2),
        "discount_amount": round(discount_amount, 2),
        "taxable": round(taxable, 2),
        "tax_amount": round(tax_amount, 2),
        "total": round(total, 2),
    }


def _fmt_money(currency: str, amount: float) -> str:
    return f"{currency}{amount:,.2f}"


def _hex_to_rgb(value: str | None, default: tuple[int, int, int]) -> tuple[int, int, int]:
    if not value:
        return default
    m = re.fullmatch(r"#?([0-9a-fA-F]{6})", value.strip())
    if not m:
        return default
    h = m.group(1)
    return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _decode_logo(data_url: str | None) -> tuple[bytes, str] | None:
    """Decode a data: URI logo into (bytes, extension). Returns None on any
    problem so a bad logo never breaks the export."""
    if not data_url or not data_url.startswith("data:"):
        return None
    try:
        header, b64 = data_url.split(",", 1)
        raw = base64.b64decode(b64)
    except Exception:  # noqa: BLE001
        return None
    ext = "png"
    if "image/jpeg" in header or "image/jpg" in header:
        ext = "jpg"
    elif "image/png" in header:
        ext = "png"
    elif "image/gif" in header:
        ext = "gif"
    return raw, ext


def _ascii(text: str) -> str:
    """Last-resort fallback for fpdf2's latin-1 core font: replace common smart
    punctuation and drop anything it can't encode. Only used when the bundled
    Unicode font could not be registered."""
    if not text:
        return ""
    repl = {
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-",
        "\u00a0": " ",
    }
    for k, v in repl.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def _register_unicode_font(pdf) -> bool:
    """Register the bundled DejaVu Sans family on this FPDF instance. Returns
    True on success (use full Unicode); False means fall back to the latin-1
    core font with ASCII sanitising so export never hard-crashes."""
    try:
        for style, fname in _FONT_FILES.items():
            path = os.path.join(_FONT_DIR, fname)
            if not os.path.isfile(path):
                return False
            pdf.add_font(_FONT_FAMILY, style, path)
        return True
    except Exception:  # noqa: BLE001 — a bad/missing font must never break export
        return False


# --- Constrained rich-text rendering -------------------------------------
# The quote agent writes proposal_text / notes using a tiny, plain-text subset:
# blank-line-separated paragraphs, '- ' bullet lines, '## ' sub-headings, and
# **bold** inline emphasis. We parse that into blocks once and render it for the
# PDF and DOCX exporters. Parsing is deliberately tolerant — anything we don't
# recognise falls through as a plain paragraph so unusual model output never
# breaks an export.

# A marker (## or -/*/•) only counts when followed by a space + real content;
# a content-less marker line (e.g. "## " or "- ") is dropped, not leaked as prose.
_BULLET_RE = re.compile(r"^[-*\u2022](?:\s+(.*\S))?\s*$")
_HEADING_RE = re.compile(r"^#{1,6}(?:\s+(.*\S))?\s*$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _strip_inline_markers(text: str) -> str:
    """Drop inline **bold** markers, leaving the literal words. Used where a
    renderer can't style runs (e.g. a heading)."""
    return _BOLD_RE.sub(r"\1", text or "")


def _inline_runs(text: str) -> list[tuple[str, bool]]:
    """Split a line into (chunk, is_bold) runs on **bold** markers. An unmatched
    trailing '**' is treated as literal text rather than dropped."""
    runs: list[tuple[str, bool]] = []
    pos = 0
    for m in _BOLD_RE.finditer(text or ""):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False))
        runs.append((m.group(1), True))
        pos = m.end()
    if text and pos < len(text):
        runs.append((text[pos:], False))
    return runs or [("", False)]


def parse_rich_blocks(text: str) -> list[dict]:
    """Parse the constrained subset into a flat list of blocks:
        {"type": "heading", "text": str}
        {"type": "bullets", "items": [str, ...]}   # consecutive bullets grouped
        {"type": "para",    "text": str}
    Blank lines separate paragraphs; unrecognised lines become prose."""
    blocks: list[dict] = []
    bullet_items: list[str] = []

    def flush_bullets() -> None:
        if bullet_items:
            blocks.append({"type": "bullets", "items": list(bullet_items)})
            bullet_items.clear()

    for raw in str(text or "").split("\n"):
        line = raw.strip()
        if not line:
            flush_bullets()
            continue
        h = _HEADING_RE.match(line)
        b = _BULLET_RE.match(line)
        if h:
            flush_bullets()
            ht = (h.group(1) or "").strip()
            if ht:
                blocks.append({"type": "heading", "text": ht})
        elif b:
            item = (b.group(1) or "").strip()
            if item:
                bullet_items.append(item)
        else:
            flush_bullets()
            blocks.append({"type": "para", "text": line})
    flush_bullets()
    return blocks


def _pdf_render_rich(
    pdf, text, *, epw, enc, set_font, set_color, unicode_ok,
    size, body_color, heading_color, line_h, para_gap,
) -> None:
    """Render the constrained subset onto the PDF: '## ' as a bold sub-heading,
    '- ' lines as bulleted items with a hanging indent, everything else as a
    body paragraph. **bold** renders via fpdf2's markdown for paras/bullets."""
    blocks = parse_rich_blocks(text)
    if not blocks:
        return
    glyph = "\u2022" if unicode_ok else "-"
    indent = 5.0
    left = pdf.l_margin
    for blk in blocks:
        btype = blk.get("type")
        if btype == "heading":
            set_font("B", size + 0.5)
            set_color(heading_color)
            pdf.set_x(left)
            pdf.multi_cell(epw, line_h, enc(_strip_inline_markers(blk["text"])), align="L")
            pdf.ln(1)
        elif btype == "bullets":
            set_font("", size)
            set_color(body_color)
            for item in blk.get("items") or []:
                pdf.set_x(left)
                pdf.cell(indent, line_h, enc(glyph))
                pdf.set_left_margin(left + indent)
                pdf.set_x(left + indent)
                try:
                    pdf.multi_cell(epw - indent, line_h, enc(item), align="L", markdown=True)
                finally:
                    pdf.set_left_margin(left)
                pdf.ln(0.6)
            pdf.ln(para_gap)
        else:
            set_font("", size)
            set_color(body_color)
            pdf.set_x(left)
            pdf.multi_cell(epw, line_h, enc(blk["text"]), align="L", markdown=True)
            pdf.ln(para_gap)


def render_pdf(quote: dict, branding: dict) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import TableBordersLayout, TableCellFillMode
    from fpdf.fonts import FontFace

    currency = quote.get("currency") or branding.get("currency") or "$"
    primary = _hex_to_rgb(branding.get("primary_color"), (24, 24, 27))
    accent = _hex_to_rgb(branding.get("accent_color"), (255, 91, 46))
    items = quote.get("line_items") or []
    totals = compute_totals(items, quote.get("tax_percent"), quote.get("discount_percent"))

    pdf = FPDF(format="A4", unit="mm")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    epw = pdf.epw

    unicode_ok = _register_unicode_font(pdf)
    fam = _FONT_FAMILY if unicode_ok else "Helvetica"
    enc = (lambda s: s) if unicode_ok else _ascii

    def font(style: str = "", size: float = 11) -> None:
        pdf.set_font(fam, style, size)

    def color(rgb: tuple[int, int, int]) -> None:
        pdf.set_text_color(*rgb)

    # --- Header: logo (left) + vendor block (right) ---
    logo = _decode_logo(branding.get("logo_data_url"))
    top = pdf.get_y()
    logo_h = 16
    if logo:
        try:
            raw, _ext = logo
            pdf.image(io.BytesIO(raw), x=pdf.l_margin, y=top, h=logo_h)
        except Exception:  # noqa: BLE001 — never let a bad logo break export
            pass
    pdf.set_xy(pdf.l_margin, top)
    font("B", 15)
    color(primary)
    vendor = branding.get("company_name") or "Proposal"
    pdf.cell(epw, 7, enc(vendor), align="R", new_x="LMARGIN", new_y="NEXT")
    font("", 8.5)
    color(_MUTED)
    for key in ("address", "email", "phone", "website"):
        val = branding.get(key)
        if val:
            pdf.cell(epw, 4.6, enc(str(val)), align="R", new_x="LMARGIN", new_y="NEXT")

    # Accent rule under the header.
    pdf.set_y(max(pdf.get_y(), top + logo_h) + 4)
    pdf.set_draw_color(*accent)
    pdf.set_line_width(0.8)
    y = pdf.get_y()
    pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
    pdf.ln(7)

    # --- Title ---
    font("B", 22)
    color(primary)
    pdf.multi_cell(epw, 9, enc(quote.get("title") or "Proposal"), align="L")
    pdf.ln(1)

    # --- Prepared for ---
    client = quote.get("client_name") or ""
    client_co = quote.get("client_company") or ""
    if client or client_co:
        font("", 10.5)
        color(_MUTED)
        who = " \u2014 ".join([p for p in (client, client_co) if p])
        pdf.cell(0, 6, enc(f"Prepared for: {who}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # --- Narrative ---
    text = quote.get("proposal_text") or ""
    if text:
        _pdf_render_rich(
            pdf, text, epw=epw, enc=enc, set_font=font, set_color=color,
            unicode_ok=unicode_ok, size=11, body_color=_BODY,
            heading_color=primary, line_h=6, para_gap=1.5,
        )
        pdf.ln(2)

    # --- Line items table ---
    if items:
        font("B", 12)
        color(primary)
        pdf.cell(0, 7, enc("Investment"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)

        headings_style = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=accent)
        font("", 10)
        color(_BODY)
        pdf.set_draw_color(228, 228, 232)
        pdf.set_line_width(0.2)
        with pdf.table(
            col_widths=(52, 12, 18, 18),
            text_align=("LEFT", "RIGHT", "RIGHT", "RIGHT"),
            borders_layout=TableBordersLayout.HORIZONTAL_LINES,
            cell_fill_color=(247, 247, 249),
            cell_fill_mode=TableCellFillMode.ROWS,
            headings_style=headings_style,
            line_height=6,
            padding=2,
        ) as table:
            head = table.row()
            for label in ("Description", "Qty", "Unit", "Amount"):
                head.cell(enc(label))
            for it in items:
                try:
                    qty = float(it.get("quantity", 0) or 0)
                    unit = float(it.get("unit_price", 0) or 0)
                except (TypeError, ValueError):
                    qty, unit = 0.0, 0.0
                row = table.row()
                row.cell(enc(str(it.get("description", "")).strip() or "\u2014"))
                row.cell(f"{qty:g}")
                row.cell(enc(_fmt_money(currency, unit)))
                row.cell(enc(_fmt_money(currency, qty * unit)))

        pdf.ln(3)

        # --- Totals block, value column aligned to the table's Amount column ---
        amt_w = epw * 0.18
        label_w = epw - amt_w

        def total_row(label: str, value: str, bold: bool = False) -> None:
            font("B" if bold else "", 11 if bold else 10)
            color(primary if bold else _MUTED)
            pdf.cell(label_w, 7, enc(label), align="R")
            color(primary if bold else _BODY)
            pdf.cell(amt_w, 7, enc(value), align="R", new_x="LMARGIN", new_y="NEXT")

        total_row("Subtotal", _fmt_money(currency, totals["subtotal"]))
        if totals["discount_amount"]:
            total_row(
                f"Discount ({quote.get('discount_percent') or 0:g}%)",
                "-" + _fmt_money(currency, totals["discount_amount"]),
            )
        if totals["tax_amount"]:
            total_row(
                f"Tax ({quote.get('tax_percent') or 0:g}%)",
                _fmt_money(currency, totals["tax_amount"]),
            )
        # Thin rule above the grand total.
        rule_x0 = max(pdf.l_margin, pdf.w - pdf.r_margin - (amt_w + 50))
        yt = pdf.get_y() + 0.5
        pdf.set_draw_color(210, 210, 214)
        pdf.set_line_width(0.3)
        pdf.line(rule_x0, yt, pdf.w - pdf.r_margin, yt)
        pdf.ln(2)
        total_row("Total", _fmt_money(currency, totals["total"]), bold=True)

    # --- Notes ---
    notes = quote.get("notes")
    if notes:
        pdf.ln(5)
        font("B", 11)
        color(primary)
        pdf.cell(0, 6, enc("Notes & Terms"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        _pdf_render_rich(
            pdf, str(notes), epw=epw, enc=enc, set_font=font, set_color=color,
            unicode_ok=unicode_ok, size=9.5, body_color=(90, 90, 95),
            heading_color=primary, line_h=5.2, para_gap=0.5,
        )

    # --- Terms & Conditions (brand boilerplate, when present) ---
    terms = branding.get("terms")
    if terms:
        pdf.ln(5)
        font("B", 11)
        color(primary)
        pdf.cell(0, 6, enc("Terms & Conditions"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        _pdf_render_rich(
            pdf, str(terms), epw=epw, enc=enc, set_font=font, set_color=color,
            unicode_ok=unicode_ok, size=9, body_color=(90, 90, 95),
            heading_color=primary, line_h=5, para_gap=0.5,
        )

    # --- Signature block ---
    signature = branding.get("signature")
    if signature:
        pdf.ln(6)
        font("", 9.5)
        color(_BODY)
        for ln in str(signature).splitlines():
            pdf.multi_cell(epw, 5, enc(ln.strip()), align="L")

    # --- Footer line ---
    footer = branding.get("footer")
    if footer:
        pdf.ln(5)
        font("", 8)
        color(_MUTED)
        for ln in str(footer).splitlines():
            pdf.multi_cell(epw, 4.2, enc(ln.strip()), align="C")

    out = pdf.output()
    return bytes(out)


def _docx_shade(cell, rgb: tuple[int, int, int]) -> None:
    """Set a table cell's background fill (python-docx has no native API)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "%02X%02X%02X" % rgb)
    tc_pr.append(shd)


def _docx_bottom_border(paragraph, rgb: tuple[int, int, int], size: int = 12) -> None:
    """Draw a horizontal rule by giving a paragraph a bottom border."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "%02X%02X%02X" % rgb)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _docx_add_runs(paragraph, text, *, size, color) -> None:
    """Add **bold**-aware runs to a paragraph, sizing and colouring each run."""
    from docx.shared import Pt, RGBColor

    for chunk, bold in _inline_runs(text):
        if not chunk:
            continue
        r = paragraph.add_run(chunk)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(*color)


def _docx_bullet_paragraph(doc):
    """A native 'List Bullet' paragraph; fall back to a manual glyph if the
    template lacks that style so the export never crashes."""
    try:
        return doc.add_paragraph(style="List Bullet")
    except (KeyError, ValueError):
        p = doc.add_paragraph()
        p.add_run("\u2022  ")
        return p


def _docx_render_rich(doc, text, *, size, body_color, heading_color, space_after) -> None:
    """Render the constrained subset into a docx: '## ' as a bold sub-heading,
    '- ' lines as List Bullet paragraphs, everything else as a body paragraph,
    with **bold** inline emphasis preserved as bold runs."""
    from docx.shared import Pt, RGBColor

    for blk in parse_rich_blocks(text):
        btype = blk.get("type")
        if btype == "heading":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(_strip_inline_markers(blk["text"]))
            r.bold = True
            r.font.size = Pt(size + 1)
            r.font.color.rgb = RGBColor(*heading_color)
        elif btype == "bullets":
            for item in blk.get("items") or []:
                p = _docx_bullet_paragraph(doc)
                p.paragraph_format.space_after = Pt(2)
                _docx_add_runs(p, item, size=size, color=body_color)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(space_after)
            _docx_add_runs(p, blk["text"], size=size, color=body_color)


def render_docx(quote: dict, branding: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    currency = quote.get("currency") or branding.get("currency") or "$"
    primary = _hex_to_rgb(branding.get("primary_color"), (24, 24, 27))
    accent = _hex_to_rgb(branding.get("accent_color"), (255, 91, 46))
    items = quote.get("line_items") or []
    totals = compute_totals(items, quote.get("tax_percent"), quote.get("discount_percent"))

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(*_BODY)

    # --- Header: logo + vendor block ---
    logo = _decode_logo(branding.get("logo_data_url"))
    if logo:
        try:
            raw, _ext = logo
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.add_run().add_picture(io.BytesIO(raw), height=Inches(0.55))
        except Exception:  # noqa: BLE001
            pass

    vendor = branding.get("company_name")
    if vendor:
        h = doc.add_paragraph()
        h.paragraph_format.space_after = Pt(0)
        r = h.add_run(vendor)
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(*primary)
    contact_bits = [str(b) for b in (branding.get(k) for k in ("address", "email", "phone", "website")) if b]
    contact_line = doc.add_paragraph()
    if contact_bits:
        cr = contact_line.add_run("  |  ".join(contact_bits))
        cr.font.size = Pt(8.5)
        cr.font.color.rgb = RGBColor(*_MUTED)
    _docx_bottom_border(contact_line, accent, size=18)

    # --- Title ---
    title = doc.add_heading(quote.get("title") or "Proposal", level=1)
    title.paragraph_format.space_before = Pt(10)
    for r in title.runs:
        r.font.color.rgb = RGBColor(*primary)

    # --- Prepared for ---
    client = quote.get("client_name") or ""
    client_co = quote.get("client_company") or ""
    if client or client_co:
        who = " \u2014 ".join([p for p in (client, client_co) if p])
        wp = doc.add_paragraph()
        wr = wp.add_run(f"Prepared for: {who}")
        wr.font.color.rgb = RGBColor(*_MUTED)
        wr.font.size = Pt(10.5)

    # --- Narrative ---
    _docx_render_rich(
        doc, quote.get("proposal_text") or "", size=10.5, body_color=_BODY,
        heading_color=primary, space_after=6,
    )

    # --- Line items table ---
    if items:
        heading = doc.add_heading("Investment", level=2)
        for r in heading.runs:
            r.font.color.rgb = RGBColor(*primary)

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for style_name in ("Light List Accent 1", "Light Grid Accent 1", "Table Grid"):
            try:
                table.style = style_name
                break
            except KeyError:
                continue

        hdr = table.rows[0].cells
        for cell, label in zip(hdr, ("Description", "Qty", "Unit", "Amount")):
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            _docx_shade(cell, accent)
        for i in (1, 2, 3):
            hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        for it in items:
            try:
                qty = float(it.get("quantity", 0) or 0)
                unit = float(it.get("unit_price", 0) or 0)
            except (TypeError, ValueError):
                qty, unit = 0.0, 0.0
            cells = table.add_row().cells
            cells[0].text = str(it.get("description", "")).strip() or "\u2014"
            cells[1].text = f"{qty:g}"
            cells[2].text = _fmt_money(currency, unit)
            cells[3].text = _fmt_money(currency, qty * unit)
            for i in (1, 2, 3):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        widths = (Inches(3.6), Inches(0.7), Inches(1.15), Inches(1.15))
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w

        def total_para(label: str, value: str, bold: bool = False) -> None:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{label}:  {value}")
            r.bold = bold
            r.font.color.rgb = RGBColor(*primary) if bold else RGBColor(*_BODY)
            if bold:
                r.font.size = Pt(12)

        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        total_para("Subtotal", _fmt_money(currency, totals["subtotal"]))
        if totals["discount_amount"]:
            total_para(
                f"Discount ({quote.get('discount_percent') or 0:g}%)",
                "-" + _fmt_money(currency, totals["discount_amount"]),
            )
        if totals["tax_amount"]:
            total_para(
                f"Tax ({quote.get('tax_percent') or 0:g}%)",
                _fmt_money(currency, totals["tax_amount"]),
            )
        total_para("Total", _fmt_money(currency, totals["total"]), bold=True)

    # --- Notes ---
    notes = quote.get("notes")
    if notes:
        nh = doc.add_heading("Notes & Terms", level=2)
        for r in nh.runs:
            r.font.color.rgb = RGBColor(*primary)
        _docx_render_rich(
            doc, str(notes), size=9.5, body_color=(90, 90, 95),
            heading_color=primary, space_after=3,
        )

    # --- Terms & Conditions (brand boilerplate, when present) ---
    terms = branding.get("terms")
    if terms:
        th = doc.add_heading("Terms & Conditions", level=2)
        for r in th.runs:
            r.font.color.rgb = RGBColor(*primary)
        _docx_render_rich(
            doc, str(terms), size=9, body_color=(90, 90, 95),
            heading_color=primary, space_after=3,
        )

    # --- Signature block ---
    signature = branding.get("signature")
    if signature:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        for ln in str(signature).splitlines():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(ln.strip())
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(*_BODY)

    # --- Footer line ---
    footer = branding.get("footer")
    if footer:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        for ln in str(footer).splitlines():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(ln.strip())
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(*_MUTED)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
